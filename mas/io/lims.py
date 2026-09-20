"""Разбор сложной шапки ЛИМС, словарь единиц, нормализация точек отбора.

Шапка: строка 0 - имя точки (тянется вправо до следующей непустой),
строка 1 - показатель, строка 2 - единица (НЕНАДЁЖНА, не используется для расчётов),
строка 3 - число значений, строки 4+ - пары колонок [время, значение]. Шаг по колонкам - 2.
"""
from __future__ import annotations

import glob
import os
import re

import openpyxl
import pandas as pd

from mas.common import get_data_dir, load_cfg

# Единицы - фиксированный словарь в коде: строка единиц в шапке файла ненадёжна.
UNITS: dict[str, str] = {
    "Mg.Sulfur": "мг/кг", "Mass.Sulfur": "% масс.", "FlashPoint": "°C",
    "95%.T": "°C", "90%.T": "°C", "50%.T": "°C", "IBP.T": "°C", "EBP.T": "°C",
    "D15": "кг/м3", "CFPP": "°C", "CloudPoint": "°C", "PourPoint": "°C",
    "I250": "%", "I350": "%", "CetaneNumber": "-", "FilterabilityLimit.T": "°C",
}

_POINT_RE = re.compile(r"Установка\s*'([^']+)'.*?Точка отбора\s*'([^']+)'")


def _normalize_point(label: str) -> str:
    """"Гидроочистка" т.N -> HTN; "АВТ" т.N -> AVTN."""
    m = _POINT_RE.search(label)
    if not m:
        return label.strip()
    unit_name, num = m.group(1), m.group(2)
    prefix = "HT" if "идроочистк" in unit_name else "AVT"
    return f"{prefix}{num}"


def _find_file() -> str:
    cfg = load_cfg()
    matches = glob.glob(os.path.join(get_data_dir(), cfg["files"]["lims"]))
    if not matches:
        raise FileNotFoundError("Не найден файл ЛИМС по шаблону из config.files.lims")
    return sorted(matches)[0]


def read_lims() -> tuple[pd.DataFrame, pd.DataFrame, int]:
    """(длинная таблица [point, parameter, time, value], метаданные единиц, число нечисловых)."""
    wb = openpyxl.load_workbook(_find_file(), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    header_point, header_param, header_unit = rows[0], rows[1], rows[2]
    data_rows = rows[4:]

    records: list[dict] = []
    units_meta: list[dict] = []
    n_non_numeric = 0
    current_point: str | None = None
    for col in range(0, len(header_point), 2):
        if header_point[col]:
            current_point = _normalize_point(str(header_point[col]))
        param = header_param[col] if col < len(header_param) else None
        if not param or current_point is None:
            continue
        declared_unit = header_unit[col] if col < len(header_unit) else None
        units_meta.append({"point": current_point, "parameter": param,
                            "declared_unit": declared_unit, "fixed_unit": UNITS.get(param, "?")})
        for row in data_rows:
            t = row[col] if col < len(row) else None
            val = row[col + 1] if col + 1 < len(row) else None
            if t is None or val is None:
                continue
            if not isinstance(val, (int, float)):
                n_non_numeric += 1
                continue
            records.append({"point": current_point, "parameter": param, "time": t, "value": float(val)})

    long_df = pd.DataFrame.from_records(records)
    units_df = pd.DataFrame.from_records(units_meta)
    return long_df, units_df, n_non_numeric
