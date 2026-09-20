"""Справочник тегов (две версии) и формулы ВАК.

Уточнённая версия приоритетна; если её нет в пакете - разбирается исходный файл.
Версия источника идёт в отчёт о данных.
"""
from __future__ import annotations

import glob
import os
import re

import openpyxl

from mas.common import get_data_dir, load_cfg

_NAME_RE = re.compile(r"^[A-Za-z0-9\-]+:[^:]+:[^:]+$")

def _glob_one(pattern: str) -> str | None:
    matches = glob.glob(os.path.join(get_data_dir(), pattern))
    return sorted(matches)[0] if matches else None

def _read_v2_dict_sheet(ws) -> dict[str, tuple[str, str | None]]:
    """Листы «АВТ» и «24-2000» уточнённого справочника: колонки «Колонка[ CSV]», «Описание[ (КИП)]», «Физическая величина»."""
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h).strip() if h else "" for h in rows[0]]
    i_tag = next(i for i, h in enumerate(header) if h.startswith("Колонка"))
    i_desc = next(i for i, h in enumerate(header) if h.startswith("Описание"))
    i_phys = header.index("Физическая величина")
    out = {}
    for row in rows[1:]:
        if row[i_tag]:
            out[str(row[i_tag]).strip()] = (row[i_desc], row[i_phys])
    return out

def _read_v1_kip_sheet(ws) -> dict[str, tuple[str, None]]:
    """Исходный формат, колонка АВТ: пары [описание, тег] - описания АВТ проверены и надёжны."""
    avt: dict[str, tuple[str, None]] = {}
    for row in list(ws.iter_rows(values_only=True))[1:]:
        if row[0] and row[1]:
            avt[str(row[1]).strip()] = (row[0], None)
    return avt

# Запасной вариант, если уточнённого справочника в пакете нет. Колонка 24-2000 исходного
# справочника сдвинута построчно (например, T11 там описан как «расход сырья»), поэтому
# 26 описаний реконструированы по проверенной таблице методических материалов.
HO_TAGS_RECONSTRUCTED: dict[str, tuple[str, str]] = {
    "F1": ("Расход бензина с установки, объёмный", "м3ч"), "F2": ("Газовая схема, расход газа от ЦК-201", "нм3ч"),
    "P3": ("Полисеп, сепаратор С-201, давление на входе", "давление"), "W4": ("К-206, массовый расход бензина", "тч"),
    "T5": ("Полисеп, Р-201, температура ГСС на выходе", "температура"), "T6": ("Полисеп, Р-202, температура ГСС на входе", "температура"),
    "W7": ("К-201, расход газа поддува, массовый", "тч"), "P8": ("Реактор Р-202, перепад давления", "перепад"),
    "F9": ("Расход сырья на установку, массовый", "тч"), "W10": ("Расход бензина с установки, массовый", "тч"),
    "T11": ("Трубопровод на выходе Р-202, температура", "температура"), "T12": ("К-201, температура верха", "температура"),
    "P13": ("Полисеп, Р-202, давление на входе", "давление"), "F14": ("Расход квенча в Р-202", "тч"),
    "F15": ("Расход сырья, объёмный", "м3ч"), "T16": ("Блок стабилизации, С-205, температура", "температура"),
    "F17": ("Расход ГО ДТ в цех №8, массовый", "тч"), "T18": ("ВАК температуры вспышки ГО ДТ", "температура"),
    "F19": ("К-201, расход бензина в колонну", "тч"), "Q20": ("Поточный анализатор серы в ДТ (Н-202)", "сера"),
    "Q21": ("Поточный анализатор серы в г/о ДТ", "сера"), "F22": ("К-201, расход газа поддува, объёмный", "нм3ч"),
    "T23": ("К-201, температура низа", "температура"), "P24": ("К-201, давление на выходе", "давление"),
    "F25": ("Расход свежего ВСГ с КЦА", "нм3ч"), "F26": ("Расход ГО ДТ в цех №8, объёмный", "м3ч"),
}

def read_tag_dictionary() -> tuple[dict, dict, str]:
    """({tag: (описание, величина)} АВТ, то же 24-2000, версия источника)."""
    cfg = load_cfg()
    v2_path = _glob_one(cfg["files"]["tags_v2"])
    if v2_path:
        wb = openpyxl.load_workbook(v2_path, read_only=True, data_only=True)
        return _read_v2_dict_sheet(wb["АВТ"]), _read_v2_dict_sheet(wb["24-2000"]), "v2 (уточнённый)"
    orig_path = _glob_one(cfg["files"]["tags"])
    if not orig_path:
        raise FileNotFoundError("Не найден ни уточнённый, ни исходный справочник тегов")
    wb = openpyxl.load_workbook(orig_path, read_only=True, data_only=True)
    avt = _read_v1_kip_sheet(wb["КИП"])
    return avt, dict(HO_TAGS_RECONSTRUCTED), "v1 (АВТ - лист КИП; 24-2000 - реконструкция по проверенной таблице)"

def _scan_formula_pairs(ws) -> list[tuple[str, str]]:
    """Ищет пары [имя_модели, формула]: формула - в соседней справа ячейке (исходный лист «ВАК» - пары подряд,
    уточнённый файл - колонки «№ | Модель | Формула | Пример расчёта»)."""
    pairs: list[tuple[str, str]] = []
    for row in ws.iter_rows(values_only=True):
        for col in range(len(row) - 1):
            name, formula = row[col], row[col + 1]
            if isinstance(name, str) and _NAME_RE.match(name.strip()) and formula:
                pairs.append((name.strip(), str(formula)))
    return pairs

def read_vak_formulas() -> tuple[list[tuple[str, str]], str]:
    """([(модель, формула)], версия источника)."""
    cfg = load_cfg()
    v2_path = _glob_one(cfg["files"]["vak_v2"])
    if v2_path:
        wb = openpyxl.load_workbook(v2_path, read_only=True, data_only=True)
        pairs = [p for sn in wb.sheetnames for p in _scan_formula_pairs(wb[sn])]
        return pairs, "v2 (уточнённый)"
    orig_path = _glob_one(cfg["files"]["tags"])
    if not orig_path:
        raise FileNotFoundError("Не найден ни уточнённый, ни исходный файл формул ВАК")
    wb = openpyxl.load_workbook(orig_path, read_only=True, data_only=True)
    if "ВАК" not in wb.sheetnames:
        raise FileNotFoundError("В исходном справочнике нет листа ВАК")
    return _scan_formula_pairs(wb["ВАК"]), "v1 (исходный, лист ВАК)"
